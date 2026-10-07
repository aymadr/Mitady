"""Word CV helpers, called by the Node server.

    python docx_tools.py extract <file.docx>
        -> prints {"text": "..."} (paragraphs, tables and text boxes)

    python docx_tools.py photo <file.docx> <out_base>
        -> saves the CV photo (largest embedded picture), prints {"photo": path|null}

    python docx_tools.py tailor <job.json>
        job.json = {"source", "output", "titleText", "newTitle", "summaryText", "newSummary",
                    "skillsAnchor", "skills": [], "referencePdf"}
        With "referencePdf" (the untouched CV as PDF), the result is tightened slightly until it
        has no more pages than the original, so a longer title never pushes a few lines onto
        a second page.
        -> writes a tailored copy, as PDF when "output" ends with .pdf and Word
           (or LibreOffice) is available, otherwise as .docx. Prints {"output": path}.

The layout of the original file is kept: only the text of existing runs is changed.
"""
import json
import os
import re
import shutil
import subprocess
import sys

import docx
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8")


def all_paragraphs(doc):
    """Every paragraph of the body, tables and text boxes, in document order."""
    for p_el in doc.element.body.iter(docx.oxml.ns.qn("w:p")):
        yield Paragraph(p_el, doc)


def extract(path):
    doc = docx.Document(path)
    lines, seen = [], set()
    for p in all_paragraphs(doc):
        t = p.text.strip()
        # Text boxes are stored twice (modern + VML fallback): skip exact repeats.
        key = (t, p._p.getparent().tag if p._p.getparent() is not None else "")
        if t and key not in seen:
            seen.add(key)
            lines.append(t)
    return {"text": "\n".join(lines)}


def photo(path, out_base):
    """Saves the largest picture embedded in the document (the CV photo) as
    <out_base>.<ext>. Prints {"photo": path} or {"photo": null}."""
    import zipfile

    with zipfile.ZipFile(path) as z:
        media = [i for i in z.infolist() if i.filename.startswith("word/media/")
                 and os.path.splitext(i.filename)[1].lower() in (".png", ".jpg", ".jpeg")]
        if not media:
            return {"photo": None}
        best = max(media, key=lambda i: i.file_size)
        out = out_base + os.path.splitext(best.filename)[1].lower()
        with open(out, "wb") as f:
            f.write(z.read(best))
        return {"photo": out}


def replace_in_paragraph(p, old, new):
    """Replaces `old` by `new` in a paragraph, keeping the formatting of the first run."""
    if old not in p.text:
        return False
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    # The text spans several runs: rebuild it inside the runs that hold it.
    full = p.text
    start = full.index(old)
    end = start + len(old)
    pos = 0
    written = False
    for r in p.runs:
        r_start, r_end = pos, pos + len(r.text)
        pos = r_end
        if r_end <= start or r_start >= end:
            continue
        before = r.text[: max(0, start - r_start)]
        after = r.text[max(0, end - r_start):] if r_end > end else ""
        r.text = before + ("" if written else new) + after
        written = True
    return True


def list_end_after(doc, p):
    """Last paragraph of the list that follows a heading: the consecutive,
    non-empty sibling paragraphs sharing the first one's style."""
    qn_p = docx.oxml.ns.qn("w:p")
    el = p._p.getnext()
    while el is not None and not (el.tag == qn_p and Paragraph(el, doc).text.strip()):
        el = el.getnext()
    if el is None:
        return None
    last = Paragraph(el, doc)
    style = last.style.style_id if last.style is not None else None
    nxt = el.getnext()
    while nxt is not None and nxt.tag == qn_p:
        q = Paragraph(nxt, doc)
        if not q.text.strip() or (q.style.style_id if q.style is not None else None) != style:
            break
        last = q
        nxt = nxt.getnext()
    return last


def pdf_pages(path):
    """Number of pages of a PDF: exact with pypdf when installed, otherwise a rough count. 0 if unknown."""
    try:
        import pypdf  # noqa: WPS433

        return len(pypdf.PdfReader(path).pages)
    except Exception:
        pass
    try:
        with open(path, "rb") as f:
            return len(re.findall(rb"/Type\s*/Page(?![s\w])", f.read()))
    except Exception:
        return 0


def _scaled(length, k):
    return None if length is None else Pt(round(length.pt * k * 2) / 2)


def tighten(doc, level):
    """Shrinks fonts by up to ~10 %, paragraph gaps and top/bottom margins a little more with each
    level, so a CV that overflows by a few lines goes back to its original page count."""
    font = max(0.9, 1 - 0.018 * level)
    gap = max(0.15, 1 - 0.2 * level)
    for st in doc.styles:
        try:
            if st.font is not None and st.font.size:
                st.font.size = _scaled(st.font.size, font)
            pf = getattr(st, "paragraph_format", None)
            if pf is not None:
                if pf.space_before:
                    pf.space_before = _scaled(pf.space_before, gap)
                if pf.space_after:
                    pf.space_after = _scaled(pf.space_after, gap)
        except Exception:
            continue
    for p in all_paragraphs(doc):
        for r in p.runs:
            if r.font.size:
                r.font.size = _scaled(r.font.size, font)
        pf = p.paragraph_format
        if pf.space_before:
            pf.space_before = _scaled(pf.space_before, gap)
        if pf.space_after:
            pf.space_after = _scaled(pf.space_after, gap)
        if isinstance(pf.line_spacing, float) and pf.line_spacing > 1.0:
            pf.line_spacing = max(1.0, round(pf.line_spacing - 0.02 * level, 2))
    for sec in doc.sections:
        for side in ("top_margin", "bottom_margin"):
            cur = getattr(sec, side)
            if cur:
                setattr(sec, side, max(Cm(0.8), cur - Cm(0.08) * level))


def fit_pages(base_docx, out_pdf, target):
    """PDF of base_docx, tightened step by step until it has at most `target` pages.
    Returns (pdf path or None, tightening level used)."""
    pdf = to_pdf(base_docx, out_pdf)
    if not pdf or not target:
        return pdf, 0
    if pdf_pages(pdf) <= target:
        return pdf, 0
    for level in range(1, 7):
        d = docx.Document(base_docx)
        tighten(d, level)
        tmp_docx = out_pdf[:-4] + ".fit.docx"
        d.save(tmp_docx)
        tmp_pdf = to_pdf(tmp_docx, out_pdf[:-4] + ".fit.pdf")
        os.remove(tmp_docx)
        if not tmp_pdf:
            break
        if pdf_pages(tmp_pdf) <= target:
            shutil.move(tmp_pdf, out_pdf)
            return out_pdf, level
        os.remove(tmp_pdf)
    return pdf, 0


def tailor(job):
    doc = docx.Document(job["source"])
    title_text = (job.get("titleText") or "").strip()
    new_title = (job.get("newTitle") or "").strip()
    anchor = (job.get("skillsAnchor") or "").strip().lower()
    skills = [s.strip() for s in job.get("skills", []) if s.strip()]

    paragraphs = list(all_paragraphs(doc))
    replaced = 0
    if title_text and new_title:
        for p in paragraphs:
            if replace_in_paragraph(p, title_text, new_title):
                replaced += 1

    # Profile paragraph: replaced only when it was rewritten for this offer.
    summary_text = (job.get("summaryText") or "").strip()
    new_summary = (job.get("newSummary") or "").strip()
    summary_replaced = False
    if summary_text and new_summary and new_summary != summary_text:
        for p in paragraphs:
            if p.text.strip() == summary_text and p.runs:
                runs = p.runs
                runs[0].text = new_summary
                for r in runs[1:]:
                    r.text = ""
                summary_replaced = True

    added = []
    if anchor and skills:
        full_text = "\n".join(p.text for p in paragraphs).lower()
        missing = [s for s in skills if s.lower() not in full_text]
        for p in paragraphs:
            t = p.text.strip().lower()
            if not t.startswith(anchor):
                continue
            # "Anchor : a, b" on one line, or a heading followed by the list.
            # The list is searched among the heading's siblings only: text boxes
            # anchored nearby would otherwise come next in document order.
            target = p if len(t) > len(anchor) + 2 else list_end_after(doc, p)
            if target is not None and missing:
                run = target.runs[-1] if target.runs else target.add_run("")
                run.text = run.text.rstrip() + ", " + ", ".join(missing)
                added = missing
            break

    out = job["output"]
    want_pdf = out.lower().endswith(".pdf")
    docx_path = out[:-4] + ".docx" if want_pdf else out
    doc.save(docx_path)

    result = {"output": docx_path, "titleReplaced": replaced, "summaryReplaced": summary_replaced, "skillsAdded": added}
    if want_pdf:
        ref = job.get("referencePdf")
        target = pdf_pages(ref) if ref and os.path.exists(ref) else 0
        pdf, level = fit_pages(docx_path, out, target)
        if pdf:
            os.remove(docx_path)
            result["output"] = pdf
            result["pages"] = pdf_pages(pdf)
            result["tightened"] = level
        else:
            result["warning"] = "Conversion PDF indisponible (Word ou LibreOffice requis) : fichier .docx conservé."
    return result


def to_pdf(src, out):
    try:
        import win32com.client  # noqa: WPS433

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        try:
            d = word.Documents.Open(os.path.abspath(src), ReadOnly=True)
            d.SaveAs(os.path.abspath(out), FileFormat=17)
            d.Close(False)
        finally:
            word.Quit()
        return out
    except Exception:
        pass
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(out), src],
            check=False, capture_output=True, timeout=120,
        )
        produced = os.path.splitext(src)[0] + ".pdf"
        if os.path.exists(produced):
            if produced != out:
                shutil.move(produced, out)
            return out
    return None


if __name__ == "__main__":
    try:
        cmd, arg = sys.argv[1], sys.argv[2]
        if cmd == "extract":
            print(json.dumps(extract(arg), ensure_ascii=False))
        elif cmd == "photo":
            print(json.dumps(photo(arg, sys.argv[3]), ensure_ascii=False))
        elif cmd == "topdf":
            pdf = to_pdf(arg, sys.argv[3])
            if not pdf:
                raise RuntimeError("Conversion PDF indisponible : installe Microsoft Word ou LibreOffice.")
            print(json.dumps({"output": pdf}))
        elif cmd == "tailor":
            with open(arg, encoding="utf-8") as f:
                print(json.dumps(tailor(json.load(f)), ensure_ascii=False))
        else:
            raise ValueError("commande inconnue : " + cmd)
    except Exception as e:  # reported to the UI as-is
        print(json.dumps({"error": str(e)}, ensure_ascii=False))
        sys.exit(1)
